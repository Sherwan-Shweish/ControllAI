"""
Exporter – formatierte Excel-Arbeitsmappe und Power-BI-Paket.

Zwei Ausgabeformate für die Weiterverarbeitung im Controlling:

    Excel        eine Arbeitsmappe mit Übersichtsblatt und je einem Blatt für
                 Plan-Ist (Umsatz, Kosten), Deckungsbeiträge, Ausreißer,
                 Transaktionen, BSC-Klassifikation und Bereinigungsprotokoll –
                 mit Zahlenformaten, Filtern, fixierter Kopfzeile und Ampelfarben
    Power BI     ein ZIP-Archiv mit CSV-Dateien im Sternschema (Fakten- und
                 Dimensionstabellen inkl. Datumstabelle) und einer Anleitung
                 für Beziehungen und erste DAX-Measures

Beide Funktionen liefern ``bytes`` und eignen sich damit direkt für
Download-Buttons im Dashboard.

Autor: Sherwan Shweish
Lizenz: MIT
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from src.analytics import (
    BereinigungsReport,
    deckungsbeitragsrechnung,
    lade_tabelle,
    plan_ist_abweichung,
    umsatz_plan_ist,
    yoy_entwicklung,
)

__all__ = [
    "CSV_ENCODING",
    "CSV_FORMAT",
    "bsc_tabelle",
    "erstelle_excel_report",
    "erstelle_powerbi_paket",
    "sammle_controlling_tabellen",
]

KOPF_FUELLUNG = PatternFill("solid", fgColor="1F3A5F")
KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")
AMPEL_FUELLUNG = {
    "rot": PatternFill("solid", fgColor="F8D7DA"),
    "gelb": PatternFill("solid", fgColor="FFF3CD"),
    "grün": PatternFill("solid", fgColor="D4EDDA"),
}
FORMAT_EUR = '#,##0.00 "€"'
FORMAT_PCT = '0.00 "%"'
FORMAT_ANTEIL = "0%"

# CSV für Power BI Desktop mit deutschem Gebietsschema; BOM, damit auch Excel Umlaute korrekt anzeigt
CSV_FORMAT: dict[str, Any] = {"sep": ";", "decimal": ",", "index": False}
CSV_ENCODING = "utf-8-sig"

BLATT_BESCHREIBUNG = {
    "Plan-Ist Umsatz": "Plan-Ist-Vergleich je Geschäftsjahr, Quartal und Produktkategorie (ohne markierte Ausreißer)",
    "Plan-Ist Kosten": "Plan-Ist-Vergleich je Kostenstelle, Kostenart und Quartal",
    "Deckungsbeitrag": "DB I und DB II je Quartal und Produktkategorie mit Vorjahresvergleich",
    "Ausreißer (Prüfung)": "Statistisch auffällige Buchungen – zur Prüfung markiert, nicht gelöscht",
    "Transaktionen": "Bereinigte Fakturapositionen",
    "BSC-Klassifikation": "KI-Vorschlag und menschliche Entscheidung je Digitalisierungsprojekt",
    "Bereinigungsprotokoll": "Eingriffe der Datenbereinigung des SAP-Rohexports",
}


# =============================================================================
# Tabellen zusammenstellen
# =============================================================================


def sammle_controlling_tabellen(
    db_path: str | Path,
    *,
    geschaeftsjahr: int | None = None,
    bsc: pd.DataFrame | None = None,
    bereinigung: BereinigungsReport | None = None,
) -> dict[str, pd.DataFrame]:
    """Stellt alle Export-Tabellen aus der Datenbank zusammen.

    Args:
        db_path: Pfad zur Controlling-Datenbank.
        geschaeftsjahr: optionaler Filter; ``None`` exportiert alle Jahre.
        bsc: Ergebnis der BSC-Klassifikation inkl. menschlicher Entscheidung.
        bereinigung: Protokoll der Datenbereinigung.
    """
    tx = lade_tabelle(db_path, "umsatz_transaktionen")
    plan = lade_tabelle(db_path, "umsatzplanung")
    kategorien = lade_tabelle(db_path, "produktkategorien")
    kosten = lade_tabelle(db_path, "plan_ist_finanzen")
    kostenstellen = lade_tabelle(db_path, "kostenstellen")
    steuerung = tx[tx["ist_ausreisser"] == 0]

    plan_ist_umsatz = umsatz_plan_ist(steuerung, plan)
    plan_ist_kosten = plan_ist_abweichung(
        kosten.merge(kostenstellen[["kostenstelle", "bezeichnung"]], on="kostenstelle", how="left"), art="kosten"
    )
    db = yoy_entwicklung(deckungsbeitragsrechnung(steuerung, kategorien), "db_ii_eur").rename(
        columns={"vorjahr_wert": "db_ii_vorjahr_eur", "yoy_abweichung": "db_ii_yoy_eur", "yoy_pct": "db_ii_yoy_pct"}
    )
    ausreisser = tx[tx["ist_ausreisser"] == 1]

    tabellen = {
        "Plan-Ist Umsatz": plan_ist_umsatz,
        "Plan-Ist Kosten": plan_ist_kosten,
        "Deckungsbeitrag": db,
        "Ausreißer (Prüfung)": ausreisser,
        "Transaktionen": tx,
    }
    if geschaeftsjahr is not None:
        tabellen = {name: df[df["geschaeftsjahr"] == geschaeftsjahr] for name, df in tabellen.items()}
    if bsc is not None and not bsc.empty:
        tabellen["BSC-Klassifikation"] = bsc
    if bereinigung is not None:
        tabellen["Bereinigungsprotokoll"] = pd.DataFrame(
            [{"kennzahl": k, "wert": str(v) if isinstance(v, list) else v}
             for k, v in bereinigung.model_dump().items()]
        )
    return {name: df.reset_index(drop=True) for name, df in tabellen.items()}


def bsc_tabelle(initiativen: list[Any], entscheidungen: dict[str, Any] | None = None) -> pd.DataFrame:
    """Überführt BSC-Vorschläge in eine Tabelle nach dem Drei-Schichten-Prinzip.

    Schicht 1: Originalzitat · Schicht 2: KI-Vorschlag (unverändert) ·
    Schicht 3: menschliche Entscheidung mit Status bestätigt/überstimmt/offen.
    """
    zeilen = []
    entscheidungen = entscheidungen or {}
    for i, item in enumerate(initiativen):
        daten = item.model_dump() if hasattr(item, "model_dump") else dict(item)
        e = entscheidungen.get(daten.get("titel", str(i)), {})
        final = e.get("finale_perspektive") or daten["perspektive"]
        geprueft = bool(e.get("geprueft", False))
        zeilen.append(
            {
                "quelle": daten.get("quelle", ""),
                "titel": daten["titel"],
                "belegzitat": daten["belegzitat"],
                "zitat_verifiziert": daten.get("zitat_verifiziert"),
                "ki_perspektive": daten["perspektive"],
                "ki_technologiefeld": daten["technologiefeld"],
                "ki_konfidenz": daten["konfidenz"],
                "ki_begruendung": daten["begruendung"],
                "pruefung_erforderlich": daten.get("pruefung_erforderlich"),
                "finale_perspektive": final,
                "geprueft": geprueft,
                "status": ("überstimmt" if final != daten["perspektive"] else "bestätigt") if geprueft else "offen",
            }
        )
    return pd.DataFrame(zeilen)


# =============================================================================
# Excel
# =============================================================================


def _blattname(name: str, vergeben: set[str]) -> str:
    """Gültiger, eindeutiger Excel-Blattname (max. 31 Zeichen, ohne : \\ / ? * [ ])."""
    basis = re.sub(r"[:\\/?*\[\]]", "-", name).strip("'")[:31] or "Blatt"
    kandidat, n = basis, 2
    while kandidat.lower() in {v.lower() for v in vergeben}:
        suffix = f" ({n})"
        kandidat = basis[: 31 - len(suffix)] + suffix
        n += 1
    vergeben.add(kandidat)
    return kandidat


def _formatiere_blatt(ws: Worksheet, df: pd.DataFrame) -> None:
    """Kopfzeile, Spaltenbreiten, Zahlenformate, Filter, fixierte Kopfzeile, Ampelfarben."""
    for zelle in ws[1]:
        zelle.fill = KOPF_FUELLUNG
        zelle.font = KOPF_SCHRIFT
        zelle.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions

    for idx, spalte in enumerate(df.columns, start=1):
        buchstabe = get_column_letter(idx)
        werte = df[spalte].astype(str).head(500)
        breite = max([len(str(spalte)), *werte.str.len().tolist()]) if len(werte) else len(str(spalte))
        ws.column_dimensions[buchstabe].width = min(max(breite + 2, 10), 60)

        name = str(spalte).lower()
        if name.endswith("_eur"):
            fmt = FORMAT_EUR
        elif name.endswith("_pct"):
            fmt = FORMAT_PCT
        elif "konfidenz" in name:
            fmt = FORMAT_ANTEIL
        else:
            fmt = None
        if fmt:
            for (zelle,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                zelle.number_format = fmt
        if name in ("belegzitat", "ki_begruendung", "beschreibung"):
            ws.column_dimensions[buchstabe].width = 60
            for (zelle,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                zelle.alignment = Alignment(wrap_text=True, vertical="top")
        if name in ("ampel", "status"):
            farben = {**AMPEL_FUELLUNG, "überstimmt": AMPEL_FUELLUNG["gelb"], "bestätigt": AMPEL_FUELLUNG["grün"]}
            for (zelle,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                if zelle.value in farben:
                    zelle.fill = farben[zelle.value]


def erstelle_excel_report(
    tabellen: dict[str, pd.DataFrame],
    *,
    titel: str = "ControllAI – Controlling-Report",
    metadaten: dict[str, Any] | None = None,
) -> bytes:
    """Schreibt alle Tabellen in eine formatierte Excel-Arbeitsmappe.

    Returns:
        Inhalt der ``.xlsx``-Datei als Bytes.
    """
    puffer = io.BytesIO()
    vergeben: set[str] = {"Übersicht"}
    namen = {name: _blattname(name, vergeben) for name in tabellen}

    with pd.ExcelWriter(puffer, engine="openpyxl") as writer:
        uebersicht = [
            ["Erstellt am", datetime.now().strftime("%d.%m.%Y %H:%M")],
            *[[k, str(v)] for k, v in (metadaten or {}).items()],
            ["Hinweis", "Synthetische Demo-Daten der fiktiven ControllAI Demo GmbH."],
            [],
            ["Tabellenblatt", "Inhalt", "Zeilen"],
            *[[namen[n], BLATT_BESCHREIBUNG.get(n, ""), len(df)] for n, df in tabellen.items()],
        ]
        pd.DataFrame(uebersicht).to_excel(writer, sheet_name="Übersicht", index=False, header=False, startrow=2)
        ws = writer.sheets["Übersicht"]
        ws["A1"] = titel
        ws["A1"].font = Font(bold=True, size=14, color="1F3A5F")
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 90
        kopfzeile = 3 + uebersicht.index(["Tabellenblatt", "Inhalt", "Zeilen"])  # startrow=2 → erste Zeile ist 3
        for zelle in ws[kopfzeile]:
            zelle.font = Font(bold=True)

        for name, df in tabellen.items():
            df = df.copy()
            for spalte in df.columns:  # Excel kennt kein pandas-NA in Objektspalten
                if df[spalte].dtype == object or str(df[spalte].dtype) in ("string", "str"):
                    df[spalte] = df[spalte].astype(object).where(df[spalte].notna(), None)
            df.to_excel(writer, sheet_name=namen[name], index=False)
            _formatiere_blatt(writer.sheets[namen[name]], df)
    return puffer.getvalue()


# =============================================================================
# Power BI
# =============================================================================

POWERBI_ANLEITUNG = """\
ControllAI – Power-BI-Paket (Sternschema)
==========================================

Import in Power BI Desktop: Daten abrufen → Text/CSV, Trennzeichen Semikolon,
Dateiursprung UTF-8, Gebietsschema Deutsch (Dezimalkomma).

Faktentabellen
  fakt_umsatz        eine Zeile je Fakturaposition (Ist)
  fakt_umsatzplan    Umsatzplanung je Quartal und Produktkategorie
  fakt_kosten        Plan und Ist je Kostenstelle, Kostenart und Quartal

Dimensionstabellen
  dim_datum          Kalender (Tag) mit Jahr, Quartal, Monat, Periode
  dim_periode        Geschäftsjahr und Quartal (Schlüssel periode_key, z. B. 2025-Q3)
  dim_produktkategorie, dim_material, dim_kunde, dim_kostenstelle

Beziehungen (1 : n, einfache Filterrichtung)
  dim_datum[datum]                          → fakt_umsatz[buchungsdatum]
  dim_periode[periode_key]                  → fakt_umsatz[periode_key]
  dim_periode[periode_key]                  → fakt_umsatzplan[periode_key]
  dim_periode[periode_key]                  → fakt_kosten[periode_key]
  dim_produktkategorie[produktkategorie]    → fakt_umsatz, fakt_umsatzplan
  dim_material[material_nr]                 → fakt_umsatz[material_nr]
  dim_kunde[kunden_nr]                      → fakt_umsatz[kunden_nr]
  dim_kostenstelle[kostenstelle]            → fakt_kosten[kostenstelle]

Vorschläge für DAX-Measures
  Umsatz Ist        = CALCULATE(SUM(fakt_umsatz[umsatz_eur]), fakt_umsatz[ist_ausreisser] = 0)
  Umsatz Plan       = SUM(fakt_umsatzplan[plan_umsatz_eur])
  Abweichung %      = DIVIDE([Umsatz Ist] - [Umsatz Plan], [Umsatz Plan])
  DB I              = CALCULATE(SUM(fakt_umsatz[deckungsbeitrag_eur]), fakt_umsatz[ist_ausreisser] = 0)
  Kosten Abw. %     = DIVIDE(SUM(fakt_kosten[ist_eur]) - SUM(fakt_kosten[plan_eur]), SUM(fakt_kosten[plan_eur]))

Markierte Ausreißer (ist_ausreisser = 1) sind enthalten, damit sie im Bericht
sichtbar bleiben; die Measures oben nehmen sie aus der Steuerungssicht heraus.
"""


def _datumstabelle(von: pd.Timestamp, bis: pd.Timestamp) -> pd.DataFrame:
    tage = pd.date_range(von.normalize().replace(month=1, day=1), bis.normalize().replace(month=12, day=31))
    monate = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober",
              "November", "Dezember"]
    return pd.DataFrame(
        {
            "datum": tage.strftime("%Y-%m-%d"),
            "jahr": tage.year,
            "quartal": "Q" + tage.quarter.astype(str),
            "periode_key": tage.year.astype(str) + "-Q" + tage.quarter.astype(str),
            "monat": tage.month,
            "monatsname": [monate[m - 1] for m in tage.month],
            "kalenderwoche": tage.isocalendar().week.to_numpy(),
            "ist_werktag": (tage.dayofweek < 5).astype(int),
        }
    )


def erstelle_powerbi_paket(db_path: str | Path, *, bsc: pd.DataFrame | None = None) -> bytes:
    """Erzeugt ein ZIP-Archiv mit Power-BI-fertigen CSV-Dateien im Sternschema."""

    def periode(df: pd.DataFrame) -> pd.Series:
        return df["geschaeftsjahr"].astype(str) + "-" + df["quartal"]

    tx = lade_tabelle(db_path, "umsatz_transaktionen")
    tx.insert(0, "periode_key", periode(tx))
    plan = lade_tabelle(db_path, "umsatzplanung")
    plan.insert(0, "periode_key", periode(plan))
    kosten = lade_tabelle(db_path, "plan_ist_finanzen")
    kosten.insert(0, "periode_key", periode(kosten))
    perioden = (
        pd.concat([tx, plan, kosten])[["periode_key", "geschaeftsjahr", "quartal"]]
        .drop_duplicates()
        .sort_values("periode_key")
    )
    daten = pd.to_datetime(tx["buchungsdatum"])

    dateien: dict[str, pd.DataFrame] = {
        "fakt_umsatz.csv": tx,
        "fakt_umsatzplan.csv": plan,
        "fakt_kosten.csv": kosten,
        "dim_datum.csv": _datumstabelle(daten.min(), daten.max()),
        "dim_periode.csv": perioden,
        "dim_produktkategorie.csv": lade_tabelle(db_path, "produktkategorien"),
        "dim_material.csv": lade_tabelle(db_path, "materialien"),
        "dim_kunde.csv": lade_tabelle(db_path, "kunden"),
        "dim_kostenstelle.csv": lade_tabelle(db_path, "kostenstellen"),
    }
    if bsc is not None and not bsc.empty:
        dateien["bsc_klassifikation.csv"] = bsc

    puffer = io.BytesIO()
    with zipfile.ZipFile(puffer, "w", compression=zipfile.ZIP_DEFLATED) as archiv:
        for name, df in dateien.items():
            archiv.writestr(name, df.to_csv(**CSV_FORMAT).encode(CSV_ENCODING))
        archiv.writestr("LIESMICH_PowerBI.txt", POWERBI_ANLEITUNG.encode("utf-8"))
    return puffer.getvalue()
