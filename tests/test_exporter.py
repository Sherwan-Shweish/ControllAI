"""Tests für src/exporter.py (Excel, Power BI) und src/textextraktion.py (PDF, Suchdiktionär)."""

from __future__ import annotations

import io
import sqlite3
import zipfile

import openpyxl
import pandas as pd
import pytest

from src.exporter import (
    CSV_ENCODING,
    CSV_FORMAT,
    _blattname,
    bsc_tabelle,
    erstelle_excel_report,
    erstelle_powerbi_paket,
    sammle_controlling_tabellen,
)
from src.seed_data import PROJEKTTEXTE
from src.textextraktion import (
    RECALL_FELD,
    Seite,
    abdeckung,
    extrahiere_pdf,
    fundstellen_text,
    ist_zahlenwerk,
    text_zu_seiten,
    vorfiltern,
)

# =============================================================================
# Excel
# =============================================================================


@pytest.fixture(scope="module")
def tabellen(demo_daten) -> dict[str, pd.DataFrame]:
    return sammle_controlling_tabellen(demo_daten.db_path, geschaeftsjahr=2025, bereinigung=demo_daten.bereinigung)


@pytest.fixture(scope="module")
def arbeitsmappe(tabellen) -> openpyxl.Workbook:
    return openpyxl.load_workbook(io.BytesIO(erstelle_excel_report(tabellen, metadaten={"Geschäftsjahr": 2025})))


def test_excel_enthaelt_alle_blaetter(arbeitsmappe) -> None:
    assert arbeitsmappe.sheetnames == [
        "Übersicht", "Plan-Ist Umsatz", "Plan-Ist Kosten", "Deckungsbeitrag", "Ausreißer (Prüfung)",
        "Transaktionen", "Bereinigungsprotokoll",
    ]
    assert arbeitsmappe["Übersicht"]["A1"].value.startswith("ControllAI")


def test_excel_zeilenzahl_und_jahresfilter(arbeitsmappe, tabellen) -> None:
    for name, df in tabellen.items():
        assert arbeitsmappe[name].max_row == len(df) + 1, name
    assert set(tabellen["Plan-Ist Umsatz"]["geschaeftsjahr"]) == {2025}


def test_excel_formatierung(arbeitsmappe) -> None:
    ws = arbeitsmappe["Plan-Ist Umsatz"]
    kopf = {zelle.value: zelle for zelle in ws[1]}
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref is not None
    assert kopf["plan_umsatz_eur"].font.bold
    assert kopf["plan_umsatz_eur"].fill.fgColor.rgb.endswith("1F3A5F")

    spalte = {name: zelle.column for name, zelle in kopf.items()}
    assert ws.cell(row=2, column=spalte["ist_umsatz_eur"]).number_format == '#,##0.00 "€"'
    assert ws.cell(row=2, column=spalte["abweichung_pct"]).number_format == '0.00 "%"'

    rote_zellen = [z for (z,) in ws.iter_rows(min_row=2, min_col=spalte["ampel"], max_col=spalte["ampel"])
                   if z.value == "rot"]
    assert rote_zellen, "Sensorik Q3 2025 sollte rot sein"
    assert all(z.fill.fgColor.rgb.endswith("F8D7DA") for z in rote_zellen)


def test_excel_mit_bsc_entscheidungen(demo_daten) -> None:
    bsc = pd.DataFrame(
        [{"titel": "Data Academy", "ki_perspektive": "Lernen & Entwicklung", "ki_konfidenz": 0.93,
          "finale_perspektive": "Lernen & Entwicklung", "status": "bestätigt"}]
    )
    tabellen = sammle_controlling_tabellen(demo_daten.db_path, geschaeftsjahr=2025, bsc=bsc)
    ws = openpyxl.load_workbook(io.BytesIO(erstelle_excel_report(tabellen)))["BSC-Klassifikation"]
    kopf = {zelle.value: zelle.column for zelle in ws[1]}
    assert ws.cell(row=2, column=kopf["ki_konfidenz"]).number_format == "0%"
    assert ws.cell(row=2, column=kopf["status"]).fill.fgColor.rgb.endswith("D4EDDA")


@pytest.mark.parametrize(
    ("name", "erwartet"),
    [("Plan/Ist: Q3 [2025]?", "Plan-Ist- Q3 -2025--"), ("x" * 40, "x" * 31)],
)
def test_blattnamen_werden_bereinigt(name: str, erwartet: str) -> None:
    assert _blattname(name, set()) == erwartet


def test_blattnamen_sind_eindeutig() -> None:
    vergeben: set[str] = set()
    assert [_blattname("Daten", vergeben) for _ in range(3)] == ["Daten", "Daten (2)", "Daten (3)"]


# =============================================================================
# Power BI
# =============================================================================


@pytest.fixture(scope="module")
def powerbi(demo_daten) -> zipfile.ZipFile:
    return zipfile.ZipFile(io.BytesIO(erstelle_powerbi_paket(demo_daten.db_path)))


def _lies(archiv: zipfile.ZipFile, name: str) -> pd.DataFrame:
    return pd.read_csv(io.BytesIO(archiv.read(name)), sep=CSV_FORMAT["sep"], decimal=CSV_FORMAT["decimal"],
                       encoding=CSV_ENCODING, dtype={"faktura_nr": str, "material_nr": str, "kunden_nr": str,
                                                     "kostenstelle": str})


def test_powerbi_paket_enthaelt_sternschema(powerbi) -> None:
    assert set(powerbi.namelist()) == {
        "fakt_umsatz.csv", "fakt_umsatzplan.csv", "fakt_kosten.csv", "dim_datum.csv", "dim_periode.csv",
        "dim_produktkategorie.csv", "dim_material.csv", "dim_kunde.csv", "dim_kostenstelle.csv",
        "LIESMICH_PowerBI.txt",
    }
    assert powerbi.read("fakt_umsatz.csv").startswith(b"\xef\xbb\xbf")  # UTF-8 mit BOM
    assert "Beziehungen" in powerbi.read("LIESMICH_PowerBI.txt").decode("utf-8")


def test_powerbi_csv_ist_verlustfrei(powerbi, demo_daten) -> None:
    fakt = _lies(powerbi, "fakt_umsatz.csv")
    with sqlite3.connect(demo_daten.db_path) as conn:
        anzahl, summe = conn.execute("SELECT COUNT(*), SUM(umsatz_eur) FROM umsatz_transaktionen").fetchone()
    assert len(fakt) == anzahl
    assert fakt["umsatz_eur"].sum() == pytest.approx(summe, abs=0.01)  # Dezimalkomma korrekt gelesen


def test_powerbi_referenzielle_integritaet(powerbi) -> None:
    fakt = _lies(powerbi, "fakt_umsatz.csv")
    kosten = _lies(powerbi, "fakt_kosten.csv")
    dims = {name: _lies(powerbi, f"{name}.csv") for name in
            ("dim_datum", "dim_periode", "dim_material", "dim_kunde", "dim_kostenstelle", "dim_produktkategorie")}
    schluessel = {"dim_datum": "datum", "dim_periode": "periode_key", "dim_material": "material_nr",
                  "dim_kunde": "kunden_nr", "dim_kostenstelle": "kostenstelle",
                  "dim_produktkategorie": "produktkategorie"}
    for name, spalte in schluessel.items():
        assert dims[name][spalte].is_unique, name
    assert set(fakt["buchungsdatum"]) <= set(dims["dim_datum"]["datum"])
    assert set(fakt["periode_key"]) <= set(dims["dim_periode"]["periode_key"])
    assert set(fakt["material_nr"]) <= set(dims["dim_material"]["material_nr"])
    assert set(fakt["kunden_nr"]) <= set(dims["dim_kunde"]["kunden_nr"])
    assert set(fakt["produktkategorie"]) <= set(dims["dim_produktkategorie"]["produktkategorie"])
    assert set(kosten["kostenstelle"]) <= set(dims["dim_kostenstelle"]["kostenstelle"])


def test_powerbi_mit_bsc(demo_daten) -> None:
    archiv = zipfile.ZipFile(io.BytesIO(erstelle_powerbi_paket(demo_daten.db_path, bsc=pd.DataFrame({"a": [1]}))))
    assert "bsc_klassifikation.csv" in archiv.namelist()


# =============================================================================
# BSC-Tabelle (Drei-Schichten-Prinzip)
# =============================================================================


def test_bsc_tabelle_status() -> None:
    basis = {"perspektive": "Kunden", "technologiefeld": "KI", "konfidenz": 0.9, "begruendung": "x",
             "belegzitat": "Zitat"}
    vorschlaege = [{**basis, "titel": "A"}, {**basis, "titel": "B"}, {**basis, "titel": "C"}]
    df = bsc_tabelle(vorschlaege, {"A": {"geprueft": True}, "B": {"geprueft": True, "finale_perspektive": "Finanzen"}})
    assert df["status"].tolist() == ["bestätigt", "überstimmt", "offen"]
    assert df["ki_perspektive"].tolist() == ["Kunden"] * 3  # KI-Vorschlag bleibt unverändert


# =============================================================================
# Textextraktion und Suchdiktionär
# =============================================================================


def _mini_pdf(seiten: list[str]) -> bytes:
    """Erzeugt ein minimales, gültiges PDF mit einer Textzeile je Seite (ohne Zusatzbibliotheken)."""
    objekte: dict[int, bytes] = {}
    kinder = []
    naechste = 4
    for text in seiten:
        seite, inhalt = naechste, naechste + 1
        naechste += 2
        kinder.append(seite)
        strom = f"BT /F1 10 Tf 40 750 Td ({text}) Tj ET".encode("latin-1")
        objekte[seite] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R "
                          f">> >> /Contents {inhalt} 0 R >>").encode()
        objekte[inhalt] = b"<< /Length %d >>\nstream\n" % len(strom) + strom + b"\nendstream"
    objekte[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objekte[2] = f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kinder)}] /Count {len(kinder)} >>".encode()
    objekte[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"
    pdf = bytearray(b"%PDF-1.4\n")
    positionen = {}
    for nummer in sorted(objekte):
        positionen[nummer] = len(pdf)
        pdf += f"{nummer} 0 obj\n".encode() + objekte[nummer] + b"\nendobj\n"
    xref = len(pdf)
    gesamt = max(objekte) + 1
    pdf += f"xref\n0 {gesamt}\n0000000000 65535 f \n".encode()
    pdf += b"".join(f"{positionen[i]:010d} 00000 n \n".encode() for i in range(1, gesamt))
    pdf += f"trailer\n<< /Size {gesamt} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(pdf)


def test_pdf_wird_seitenweise_gelesen_und_vorgefiltert() -> None:
    pytest.importorskip("pdfplumber")
    pdf = _mini_pdf([
        "Seit 2025 steuern KI-gestuetzte Absatzprognosen die Disposition in allen Filialen des Konzerns.",
        "Umsatz 2025 2024 12.345 67.890 1.234 5.678 9.012 3.456 7.890 1.234 5.678 9.012 45,6 % 12,3 %",
        "Der Aufsichtsrat hat im Berichtsjahr fuenf ordentliche Sitzungen abgehalten und beraten.",
    ])
    seiten = extrahiere_pdf(pdf)
    assert [s.nummer for s in seiten] == [1, 2, 3]
    assert "Absatzprognosen" in seiten[0].text
    funde = vorfiltern(seiten)
    assert [f.seite for f in funde] == [1]  # Zahlenwerk und themenfremde Seite fallen heraus
    assert funde[0].treffer["KI"] == ["ki"]
    assert abdeckung(seiten, funde)["seiten_mit_treffern"] == 1


def test_beispieltexte_werden_vollstaendig_gefunden() -> None:
    seiten = text_zu_seiten("\f".join(t["text"] for t in PROJEKTTEXTE))
    funde = vorfiltern(seiten)
    projektseiten = {i + 1 for i, t in enumerate(PROJEKTTEXTE) if t["erwartung"]["ist_projekt"]}
    assert projektseiten <= {f.seite for f in funde}  # Recall 100 % auf den Beispieltexten
    assert "Data Academy" in fundstellen_text(funde)


def test_recall_feld_ist_keine_technologiekategorie() -> None:
    funde = vorfiltern([Seite(1, "Die Digitalisierung bleibt für uns eine der wichtigsten Prioritäten im Konzern.")])
    assert list(funde[0].treffer) == [RECALL_FELD]
    assert funde[0].technologiefelder == []


@pytest.mark.parametrize(
    "text",
    [
        "Unsere Kinder- und Jugendstiftung förderte im Berichtsjahr zahlreiche Bildungsprojekte in der Region.",
        "Die Skizze des neuen Verwaltungsgebäudes wurde im Rahmen eines Architekturwettbewerbs vorgestellt.",
    ],
)
def test_ki_wird_nur_als_eigenes_wort_erkannt(text: str) -> None:
    assert vorfiltern([Seite(1, text)]) == []


def test_zahlenwerk_erkennung() -> None:
    assert ist_zahlenwerk("2025 2024 12.345 67.890 1.234 45,6 %")
    assert not ist_zahlenwerk("Wir haben 2025 eine Plattform für 3.000 Kunden eingeführt.")
